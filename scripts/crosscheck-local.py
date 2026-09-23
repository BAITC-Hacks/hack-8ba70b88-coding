"""Read-only independent openpyxl check of the full import's representative cells.

Optional developer audit, not needed to start the app. Real results stay ignored.
Run after node scripts/verify-local-import.js; requires openpyxl.
"""
import collections
import datetime
import json
import os
import pathlib
import re
import openpyxl

base = pathlib.Path(__file__).resolve().parents[1]
data = json.loads((base / 'private-data/imported.json').read_text(encoding='utf-8'))
paths = {
    'systeme': pathlib.Path(os.environ.get('SYSTEME_DATA_DIR', pathlib.Path.home() / 'Downloads/Systeme electric/Systeme electric')),
    'iek': pathlib.Path(os.environ.get('IEK_DATA_DIR', pathlib.Path.home() / 'Downloads/IEK/IEK')),
}
books = {}
checks = []
def cell(supplier, source, expected, field):
    match = re.match(r'^(.+) / (.+)!([A-Z]+[0-9]+)', source or '')
    assert match, f'Unrecognized source: {source}'
    filename, sheet, address = match.groups()
    assert pathlib.Path(filename).name == filename
    bookkey = (supplier, filename)
    if bookkey not in books:
        books[bookkey] = openpyxl.load_workbook(paths[supplier] / filename, read_only=False, data_only=True)
    actual = books[bookkey][sheet][address].value
    if isinstance(actual, str) and actual.startswith('#'): actual = None
    if isinstance(actual, (int, float)) and isinstance(expected, (int, float)):
        assert abs(actual - expected) < 1e-7, (source, actual, expected)
    else:
        assert actual == expected, (source, actual, expected)
    checks.append({'supplier': supplier, 'source': source, 'field': field, 'actual': actual, 'imported': expected})

selected = []
for supplier in paths:
    products = [p for p in data['products'] if p['supplierId'] == supplier]
    candidates = [
        next(p for p in products if p['inbound'] and any(s['quantity'] is not None for s in p['stocks'])),
        next(p for p in products if p['monthlySales'] and p['detailMonthly'] and (p['packSize'] or p['moq'])),
        next(p for p in products if p['stocks'] and p['stocks'][-1]['quantity'] is None),
    ]
    for p in candidates:
        if p['id'] in {x['id'] for x in selected}: continue
        selected.append(p)
        for field, sourcefield in [('packSize', 'packSource'), ('moq', 'moqSource')]:
            if p[field] is not None and p[sourcefield]: cell(supplier, p[sourcefield], p[field], field)
        monthly = [r for r in p['monthlySales'] if r['quantity'] is not None]
        for record in (monthly[:1] + monthly[-1:]): cell(supplier, record['source'], record['quantity'], 'monthlySales')
        for record in (p['stocks'][-1:] + [r for r in p['stocks'] if r['basis'] == 'free']):
            cell(supplier, record['source'], record['quantity'], 'stock')
        for record in p['inbound']: cell(supplier, record['source'], record['quantity'], 'inbound')

for supplier in paths:
    wanted = {p['code']: p for p in selected if p['supplierId'] == supplier}
    filename = next(p for p in paths[supplier].glob('*.xlsx') if 'Динамика' in p.name)
    wb = openpyxl.load_workbook(filename, read_only=True, data_only=True)
    grouped = collections.defaultdict(lambda: {'positive': 0, 'negative': 0, 'rows': 0})
    for vals in wb.worksheets[0].iter_rows(min_row=2, values_only=True):
        if len(vals) < 8 or str(vals[3]).strip() not in wanted or not isinstance(vals[7], (int, float)): continue
        if not str(vals[2]).startswith('Расходная накладная'): continue
        dt = vals[0] if isinstance(vals[0], datetime.datetime) else datetime.datetime.strptime(str(vals[0]).split()[0], '%d.%m.%Y')
        key = (str(vals[3]).strip(), dt.strftime('%Y-%m'), 'warehouse:' + str(vals[6]).strip())
        record = grouped[key]
        record['positive' if vals[7] >= 0 else 'negative'] += vals[7]
        record['rows'] += 1
    wb.close()
    for code, p in wanted.items():
        for record in p['detailMonthly']:
            actual = grouped[(code, record['month'], record['warehouseId'])]
            for field in ['positive', 'negative', 'rows']:
                assert abs(actual[field] - record[field]) < 1e-7, (p['id'], record['month'], field, actual[field], record[field])
            checks.append({'supplier': supplier, 'productId': p['id'], 'month': record['month'], 'field': 'detailMonthly', 'matched': True})
for wb in books.values(): wb.close()
report = {'productsChecked': len(selected), 'checksPassed': len(checks), 'checks': checks}
(base / 'private-data/crosscheck.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'productsChecked': len(selected), 'checksPassed': len(checks)}, ensure_ascii=False))
